"""reportkit: Slab & Rule design system for python-docx (A4 report layouts)."""
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent.parent / "assets"
T = json.loads(Path(os.environ.get("REPORT_TOKENS", ASSETS / "tokens.json")).read_text())  # design-picker --tokens-out
C = T["color"]
SERIF, SANS = T["font"]["serif"], T["font"]["sans"]
PW, PH = T["page"]["size_mm"]
M = T["page"]["margin_mm"]
CW = PW - M["left"] - M["right"]  # content width in mm


# ---------------------------------------------------------------- low-level XML
def _rgb(h):
    return RGBColor.from_string(h)


def _set_fonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + a), name)


def _track(rpr, twentieths):
    sp = rpr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        rpr.append(sp)
    sp.set(qn("w:val"), str(int(twentieths)))


def shade(cell, hex_):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcpr.append(shd)


def borders(cell, **sides):
    """sides: top/left/bottom/right = (pt, hex) ; unspecified sides stay nil."""
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:tcBorders")):
        tcpr.remove(old)
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + side)
        if side in sides and sides[side]:
            pt, col = sides[side]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(max(2, int(round(pt * 8)))))
            el.set(qn("w:color"), col)
            el.set(qn("w:space"), "0")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcpr.append(b)


def pad(cell, top=0, left=0, bottom=0, right=0):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:tcMar")):
        tcpr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(int(round(v * 56.693))))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def valign(cell, where):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:vAlign")):
        tcpr.remove(old)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), where)
    tcpr.append(v)


def text_direction(cell, val="btLr"):
    tcpr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), val)
    tcpr.append(td)


def row_flags(row, header=False, cant_split=True):
    trpr = row._tr.get_or_add_trPr()
    if cant_split:
        trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def p_border(par, **sides):
    ppr = par._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            pt, col, space = sides[side]
            el = OxmlElement("w:" + side)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(int(round(pt * 8))))
            el.set(qn("w:space"), str(space))
            el.set(qn("w:color"), col)
            bd.append(el)
    ppr.append(bd)


def make_table(container, nrows, widths_mm):
    t = container.add_table(rows=nrows, cols=len(widths_mm))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    tw = tblpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblpr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int(sum(widths_mm) * 56.693)))
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    for i, w in enumerate(widths_mm):
        t.columns[i].width = Mm(w)
        for r in t.rows:
            r.cells[i].width = Mm(w)
    return t


def set_row_height(row, mm, exact=False):
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


# ---------------------------------------------------------------- styles
def _style(doc, name, base="Normal", font=SANS, size=10, bold=False, italic=False, color=None,
           caps=False, track=0, before=0, after=0, line=1.2, align=None, keep_next=False,
           left=None, first=None, builtin=False):
    names = [s.name for s in doc.styles]
    st = doc.styles[name] if (builtin or name in names) else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base] if base != name else None
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.italic = italic
    st.font.all_caps = caps
    st.font.color.rgb = _rgb(color or C["ink"])
    rpr = st.element.get_or_add_rPr()
    _set_fonts(rpr, font)
    if track:
        _track(rpr, track)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Mm(left)
    if first is not None:
        pf.first_line_indent = Mm(first)
    return st


def setup_styles(doc):
    tp, ld = T["type_pt"], T["leading"]
    n = doc.styles["Normal"]
    n.font.name, n.font.size = SERIF, Pt(tp["body"])
    _set_fonts(n.element.get_or_add_rPr(), SERIF)
    n.font.color.rgb = _rgb(C["ink"])
    n.paragraph_format.line_spacing = ld["body"]
    n.paragraph_format.space_after = Pt(T["space_pt"]["para_after"])
    n.paragraph_format.widow_control = True

    S = lambda *a, **k: _style(doc, *a, **k)
    S("Heading 1", font=SERIF, size=tp["h1"], line=ld["h1"], after=T["space_pt"]["h1_after"], keep_next=True, builtin=True)
    S("Heading 2", font=SERIF, size=tp["h2"], bold=True, before=14, after=4, line=1.2, keep_next=True, builtin=True)
    S("Heading 3", font=SANS, size=tp["label"], bold=True, caps=True, track=18, color=C["accent"], before=10, after=3, keep_next=True, builtin=True)
    S("Kicker", font=SANS, size=tp["label"], bold=True, caps=True, track=22, color=C["accent"], after=6, keep_next=True)
    S("Label", font=SANS, size=tp["label"], bold=True, caps=True, track=18, color=C["grey"], after=2)
    S("Lead", font=SERIF, size=tp["lead"], line=ld["lead"], after=10)
    S("Display", font=SERIF, size=tp["display"], line=1.02, after=10)
    S("Meta", font=SANS, size=8.5, line=1.3, after=0)
    S("Note", font=SANS, size=tp["note"], color=C["grey"], line=1.35, after=6)
    S("Caption", font=SANS, size=tp["caption"], color=C["grey"], line=1.3, before=3, after=10)
    S("FigLabel", font=SANS, size=tp["label"], bold=True, caps=True, track=14, after=4, keep_next=True)
    S("TblText", font=SANS, size=tp["table"], line=ld["table"])
    S("TblHead", font=SANS, size=7, bold=True, caps=True, track=12, line=1.15, color=C["ink"])
    S("Quote", font=SERIF, size=tp["quote"], italic=True, color=C["accent"], line=1.25, before=4, after=6)
    S("NumXL", font=SERIF, size=tp["numeral_xl"], color=C["accent"], line=1.0, after=0)
    S("NumL", font=SERIF, size=tp["numeral_l"], color=C["accent"], line=1.0, after=0)
    S("KpiVal", font=SERIF, size=tp["numeral_m"], line=1.0, after=2)
    S("KpiLab", font=SANS, size=tp["caption"], color=C["grey"], line=1.3, after=0)
    S("Callout", font=SERIF, size=11, line=1.4, after=4)
    S("Ref", font=SANS, size=tp["reference"], line=1.35, after=3, left=6, first=-6)
    S("Toc1", font=SERIF, size=15, line=1.2, before=14, after=3, keep_next=False)
    S("Toc2", font=SANS, size=9.5, line=1.3, before=2, after=3, left=14)
    S("Bullet", font=SERIF, size=tp["body"], line=ld["body"], after=3, left=5, first=-5)
    S("RunHead", font=SANS, size=7, caps=True, track=14, color=C["grey"], after=0, line=1.0)
    S("SlabLabel", font=SANS, size=tp["label"], bold=True, caps=True, track=18, color=C["accent_tint"], after=4)
    S("SlabNum", font=SERIF, size=30, color="FFFFFF", line=1.0, after=0)
    S("SlabText", font=SANS, size=8, color="FFFFFF", line=1.35, after=10)
    for nm in ("SlabLabel", "SlabNum", "SlabText"):
        doc.styles[nm].font.color.rgb = _rgb("FFFFFF" if nm != "SlabLabel" else C["accent_tint"])


# ---------------------------------------------------------------- inline markup
_TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[\^\d+(?:\s*,\s*\d+)*\])")


class Cites(set):
    pass


def runs(par, text, cites=None, **fmt):
    for tok in _TOK.split(text or ""):
        if not tok:
            continue
        r = None
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("[^"):
            nums = tok[2:-1]
            r = par.add_run(nums.replace(" ", ""))
            r.font.superscript = True
            _set_fonts(r._r.get_or_add_rPr(), SANS)
            if cites is not None:
                cites.update(int(x) for x in re.split(r"\s*,\s*", nums))
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            r = par.add_run(tok)
        for k, v in fmt.items():
            setattr(r.font, k, v)
    return par


def para(target, text="", style="Normal", cites=None, align=None, **pf):
    ps = getattr(target, "paragraphs", None)
    if hasattr(target, "_tc") and len(ps) == 1 and not ps[0].text and not ps[0].runs:
        p = ps[0]
        p.style = style
    else:
        p = target.add_paragraph(style=style)
    if not text:
        p.add_run("")  # marks the paragraph as taken so the next para() call does not reuse it
    runs(p, text, cites)
    if align is not None:
        p.alignment = align
    for k, v in pf.items():
        setattr(p.paragraph_format, k, v)
    return p


def tiny(par, after=0):
    """Collapse a paragraph to a 1pt line (spacers, rules, section-break carriers)."""
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(after)
    pf.line_spacing = Pt(1)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    ppr = par._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    rpr.append(sz)


def rule_para(target, container_mm, length_mm=22, after=14, before=0):
    """Short accent rule: a 1pt paragraph whose right indent leaves only length_mm of border."""
    p = para(target, "", "Normal")
    tiny(p, after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.right_indent = Mm(max(0, container_mm - length_mm))
    p_border(p, bottom=(T["rule_pt"]["heavy"], C["accent"], 1))
    return p


_ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings",
}


def fix_order(root):
    """Re-sort property children into OOXML schema order so Word opens the file cleanly."""
    for tag, order in _ORDER.items():
        rank = {qn("w:" + n): i for i, n in enumerate(order.split())}
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag, 999))
            for k in kids:
                el.remove(k)
            for k in kids:
                el.append(k)


# ---------------------------------------------------------------- charts
def render_chart(spec, width_mm, height_mm, out_png):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager as fm
    import matplotlib.pyplot as plt

    for f in (ASSETS / "fonts").glob("*.ttf"):
        fm.fontManager.addfont(str(f))
    plt.rcParams.update({"font.family": SANS, "font.size": 7.5, "text.color": "#" + C["ink"],
                         "axes.edgecolor": "#" + C["ink"], "axes.labelcolor": "#" + C["grey"],
                         "xtick.color": "#" + C["ink"], "ytick.color": "#" + C["ink"]})
    ink, acc, grey, light = "#" + C["ink"], "#" + C["accent"], "#8E9692", "#C9CFCB"
    kind = spec.get("kind", "bar")
    cats = spec["categories"]
    series = spec["series"]
    fig, ax = plt.subplots(figsize=(width_mm / 25.4, height_mm / 25.4), dpi=300)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(length=0)
    fmt = spec.get("format", "{:,.0f}")
    hi = spec.get("highlight", len(cats) - 1 if len(series) == 1 else None)
    palette = ["#" + c for c in T["chart_series"]]

    if kind == "bar":
        n = len(series)
        w = 0.7 / n
        for si, s in enumerate(series):
            xs = [i + (si - (n - 1) / 2) * w for i in range(len(cats))]
            cols = [(acc if i == hi else light) for i in range(len(cats))] if n == 1 else palette[si]
            bars = ax.bar(xs, s["values"], width=w * 0.92, color=cols)
            for b, v in zip(bars, s["values"]):
                ax.text(b.get_x() + b.get_width() / 2, b.get_height(), fmt.format(v), ha="center",
                        va="bottom", fontsize=7, color=ink, clip_on=False)
        ax.set_xticks(range(len(cats)), cats)
        ax.yaxis.grid(True, color="#" + C["hairline"], linewidth=0.5)
        ax.set_axisbelow(True)
        ax.set_yticks([])
        ax.yaxis.grid(False)
        ax.margins(y=0.15)
    elif kind == "hbar":
        s = series[0]
        order = list(range(len(cats)))[::-1]
        cols = [(acc if i == hi else light) for i in order] if hi is not None else [acc] * len(cats)
        bars = ax.barh([cats[i] for i in order], [s["values"][i] for i in order], color=cols, height=0.62)
        for b, i in zip(bars, order):
            ax.text(b.get_width(), b.get_y() + b.get_height() / 2, "  " + fmt.format(s["values"][i]),
                    va="center", fontsize=7, color=ink)
        ax.spines["bottom"].set_visible(False)
        ax.set_xticks([])
        ax.margins(x=0.12)
    elif kind == "line":
        for si, s in enumerate(series):
            col = palette[si] if len(series) > 1 else acc
            ax.plot(cats, s["values"], color=col, linewidth=1.6 if si == 0 else 1.1, marker="o", markersize=2.8)
            ax.text(len(cats) - 1, s["values"][-1], "  " + s.get("name", ""), va="center", fontsize=7, color=col)
            ax.text(0, s["values"][0], fmt.format(s["values"][0]) + "  ", va="center", ha="right", fontsize=7, color=col)
            ax.text(len(cats) - 1, s["values"][-1], "", fontsize=1)
        ax.yaxis.grid(True, color="#" + C["hairline"], linewidth=0.5)
        ax.set_axisbelow(True)
        ax.margins(x=0.14, y=0.18)
        ax.tick_params(axis="y", labelsize=7)
    else:
        raise ValueError("unknown chart kind: " + kind)
    if spec.get("unit"):
        ax.set_title(spec["unit"], loc="left", fontsize=7, color="#" + C["grey"], pad=8)
    if len(series) > 1 and kind == "bar":
        ax.legend([s.get("name", "") for s in series], frameon=False, loc="upper left",
                  ncol=len(series), fontsize=7, handlelength=0.8, bbox_to_anchor=(0, 1.12))
    fig.tight_layout(pad=0.3)
    fig.savefig(out_png, transparent=True)
    plt.close(fig)


# ---------------------------------------------------------------- report
_NUM = re.compile(r"^[\s\-\u2013\u2212+~<>$\u20ac\u00a3]*[\d][\d.,\s]*\s*(%|[A-Za-z]{0,3})?\s*[\u25b2\u25bc\u2191\u2193]?$")


class Report:
    def __init__(self, meta, outdir=".", pages=None):
        self.meta = meta
        self.outdir = Path(outdir)
        self.pages = pages or {}
        self.doc = Document()
        setup_styles(self.doc)
        self.cites = Cites()
        self.refs = set()
        self.fig_n = 0
        self.tbl_n = 0
        self.first = True
        self.sequence = []
        self.toc_entries = []
        self.doc.core_properties.title = meta.get("title", "")
        self.doc.core_properties.author = meta.get("author", "")

    # ----- sections
    def section(self, kind, running=None, margins=None, landscape=False, cols=1, frame=False, continuous=False):
        d = self.doc
        if self.first:
            sec = d.sections[0]
            self.first = False
        else:
            sec = d.add_section(WD_SECTION.CONTINUOUS if continuous else WD_SECTION.NEW_PAGE)
        w, h = (PH, PW) if landscape else (PW, PH)
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Mm(w), Mm(h)
        m = margins or M
        sec.top_margin, sec.bottom_margin = Mm(m["top"]), Mm(m["bottom"])
        sec.left_margin, sec.right_margin = Mm(m["left"]), Mm(m["right"])
        sec.header_distance, sec.footer_distance = Mm(T["page"]["header_mm"]), Mm(T["page"]["footer_mm"])
        sp = sec._sectPr
        for old in sp.findall(qn("w:pgBorders")):
            sp.remove(old)
        cl = sp.find(qn("w:cols"))
        if cl is None:
            cl = OxmlElement("w:cols")
            sp.append(cl)
        for a in list(cl.attrib):
            del cl.attrib[a]
        cl.set(qn("w:num"), str(cols))
        cl.set(qn("w:space"), str(int(8 * 56.693)))
        if frame:
            pb = OxmlElement("w:pgBorders")
            pb.set(qn("w:offsetFrom"), "page")
            for side in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + side)
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
                el.set(qn("w:space"), "28"); el.set(qn("w:color"), C["rule"])
                pb.append(el)
            sp.find(qn("w:pgMar")).addnext(pb)
        if not continuous:
            self._header_footer(sec, running, landscape)
            self.sequence.append(kind)
        return sec

    def _header_footer(self, sec, running, landscape=False):
        for part in (sec.header, sec.footer):
            part.is_linked_to_previous = False
            for p in part.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            part.paragraphs[0].text = ""
        if running is None:
            sec.header_distance = sec.footer_distance = Mm(0)
            tiny(sec.header.paragraphs[0])
            tiny(sec.footer.paragraphs[0])
            return
        width = (PH if landscape else PW) - M["left"] - M["right"]
        hp = sec.header.paragraphs[0]
        hp.style = "RunHead"
        hp.paragraph_format.tab_stops.add_tab_stop(Mm(width), WD_TAB_ALIGNMENT.RIGHT)
        hp.add_run(self.meta.get("short_title", self.meta.get("title", "")) + "\t" + running)
        p_border(hp, bottom=(T["rule_pt"]["hair"], C["rule"], 6))
        fp = sec.footer.paragraphs[0]
        fp.style = "RunHead"
        fp.paragraph_format.tab_stops.add_tab_stop(Mm(width), WD_TAB_ALIGNMENT.RIGHT)
        note = "Illustrative sample content" if self.meta.get("illustrative") else self.meta.get("footer", "")
        fp.add_run(note + "\t")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        _set_fonts(rpr, SERIF)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "18")
        col = OxmlElement("w:color"); col.set(qn("w:val"), C["ink"])
        rpr.append(col); rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = "1"
        r.append(t); fld.append(r)
        fp._p.append(fld)

    # ----- shared pieces
    def heading(self, title, kicker=None, lead=None, rule=True, width=CW):
        d = self.doc
        if kicker:
            para(d, kicker, "Kicker")
        para(d, title, "Heading 1")
        if rule:
            rule_para(d, width)
        if lead:
            lp = para(d, lead, "Lead", self.cites)
            lp.paragraph_format.right_indent = Mm(max(0, width - 135))

    def figure_chart(self, target, spec, width_mm):
        self.fig_n += 1
        h = spec.get("height_mm", 62)
        png = self.outdir / f"_chart_{self.fig_n}.png"
        render_chart(spec, width_mm, h, png)
        para(target, f"Figure {self.fig_n}  \u2014  {spec.get('title', '')}", "FigLabel")
        p = para(target, "", "Normal", space_after=Pt(0), keep_with_next=True)
        p.add_run().add_picture(str(png), width=Mm(width_mm))
        if spec.get("source") or spec.get("caption"):
            para(target, " ".join(filter(None, [spec.get("caption"), ("Source: " + spec["source"]) if spec.get("source") else ""])), "Caption", self.cites)

    def figure_image(self, target, spec, width_mm):
        self.fig_n += 1
        path = spec.get("path")
        para(target, f"Figure {self.fig_n}  \u2014  {spec.get('title', '')}", "FigLabel")
        if path and Path(path).exists():
            p = para(target, "", "Normal", space_after=Pt(0), keep_with_next=True)
            p.add_run().add_picture(str(path), width=Mm(width_mm))
        else:
            t = make_table(target, 1, [width_mm])
            c = t.cell(0, 0)
            shade(c, C["tint"])
            pad(c, 4, 4, 4, 4)
            valign(c, "center")
            set_row_height(t.rows[0], spec.get("height_mm", 50))
            para(c, "Image to be supplied", "Label", align=WD_ALIGN_PARAGRAPH.CENTER)
            para(c, spec.get("alt", spec.get("title", "")), "Note", align=WD_ALIGN_PARAGRAPH.CENTER)
            tiny(target.add_paragraph())
        if spec.get("caption") or spec.get("source"):
            para(target, " ".join(filter(None, [spec.get("caption"), ("Source: " + spec["source"]) if spec.get("source") else ""])), "Caption", self.cites)

    def data_table(self, target, spec, width_mm):
        hdr = spec["header"]
        rows = spec["rows"]
        ncol = len(hdr)
        self.tbl_n += 1
        if spec.get("title"):
            para(target, f"Table {self.tbl_n}  \u2014  {spec['title']}", "FigLabel")
        align = spec.get("align") or []
        if not align:
            for ci in range(ncol):
                vals = [str(r[ci]) for r in rows if ci < len(r) and str(r[ci]).strip()]
                align.append("right" if vals and all(_NUM.match(v) for v in vals) else "left")
        weights = spec.get("widths")
        if not weights:
            weights = []
            for ci in range(ncol):
                lens = [len(str(hdr[ci]))] + [len(str(r[ci])) for r in rows if ci < len(r)]
                lo, hi = (6, 12) if align[ci] == "right" else (8, 34)
                weights.append(min(max(max(lens), lo), hi))
        tot = sum(weights)
        widths = [width_mm * w / tot for w in weights]
        dense = ncol > 6 or len(rows) > 18
        size = T["type_pt"]["table_dense"] if dense else T["type_pt"]["table"]
        t = make_table(target, 1 + len(rows), widths)
        hc = spec.get("highlight_col")
        total = spec.get("total_row", False)
        vpad = 1.0 if dense else 1.5
        al = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}
        rule, hair = C["ink"], C["rule"]
        for ri, row in enumerate([hdr] + rows):
            is_h, is_last = ri == 0, ri == len(rows)
            row_flags(t.rows[ri], header=is_h)
            for ci in range(ncol):
                c = t.cell(ri, ci)
                txt = str(row[ci]) if ci < len(row) else ""
                pad(c, vpad + (0.4 if is_h else 0), 0 if ci == 0 else 1.5, vpad + (0.4 if is_h else 0), 0 if ci == ncol - 1 else 1.5)
                valign(c, "bottom" if is_h else "top")
                b = {}
                if is_h:
                    b = {"top": (T["rule_pt"]["base"], rule), "bottom": (T["rule_pt"]["base"], rule)}
                else:
                    b = {"bottom": ((T["rule_pt"]["base"], rule) if is_last else (T["rule_pt"]["hair"], hair))}
                    if total and is_last:
                        b["top"] = (T["rule_pt"]["base"], rule)
                borders(c, **b)
                if hc is not None and ci == hc and not is_h:
                    shade(c, C["tint"])
                p = para(c, txt, "TblHead" if is_h else "TblText", self.cites, align=al[align[ci]])
                if not is_h:
                    for r in p.runs:
                        r.font.size = Pt(size)
                        if (total and is_last) or (hc is not None and ci == hc):
                            r.bold = True
        if spec.get("source") or spec.get("note"):
            para(target, " ".join(filter(None, [spec.get("note"), ("Source: " + spec["source"]) if spec.get("source") else ""])), "Caption", self.cites)
        else:
            tiny(target.add_paragraph(), after=10)

    def kpis(self, target, items, width_mm):
        n = len(items)
        t = make_table(target, 1, [width_mm / n] * n)
        for i, k in enumerate(items):
            c = t.cell(0, i)
            borders(c, top=(T["rule_pt"]["base"], C["ink"]), left=(T["rule_pt"]["hair"], C["rule"]) if i else None)
            pad(c, 3, 4 if i else 0, 2, 3)
            p = para(c, "", "KpiVal")
            r = p.add_run(str(k["value"]))
            if k.get("unit"):
                u = p.add_run(k["unit"])
                u.font.size = Pt(14)
            if k.get("accent"):
                r.font.color.rgb = _rgb(C["accent"])
                if k.get("unit"):
                    u.font.color.rgb = _rgb(C["accent"])
            para(c, k["label"], "KpiLab", self.cites)
        spacer = target.add_paragraph()
        tiny(spacer)
        spacer.paragraph_format.space_after = Pt(10)

    def callout(self, target, label, text):
        p = para(target, label, "Label", space_before=Pt(6), keep_with_next=True)
        p.paragraph_format.left_indent = Mm(5)
        p_border(p, left=(T["rule_pt"]["heavy"], C["accent"], 8))
        for chunk in (text if isinstance(text, list) else [text]):
            q = para(target, chunk, "Callout", self.cites)
            q.paragraph_format.left_indent = Mm(5)
            p_border(q, left=(T["rule_pt"]["heavy"], C["accent"], 8))

    def flow(self, target, items, width_mm):
        """Render a list of block items into a container (document or cell)."""
        for it in items:
            if isinstance(it, str):
                it = {"p": it}
            k = next(iter(it))
            v = it[k]
            if k == "p":
                para(target, v, "Normal", self.cites)
            elif k == "lead":
                para(target, v, "Lead", self.cites)
            elif k == "h":
                para(target, v, "Heading 2", keep_with_next=not hasattr(target, "_tc"))
            elif k == "h3":
                para(target, v, "Heading 3")
            elif k == "bullets":
                for b in v:
                    para(target, "\u2013\t" + b, "Bullet", self.cites)
            elif k == "quote":
                p = para(target, v, "Quote", self.cites, space_before=Pt(10))
                p_border(p, top=(T["rule_pt"]["base"], C["accent"], 8))
                if it.get("by"):
                    para(target, it["by"], "Label", space_after=Pt(10))
            elif k == "callout":
                self.callout(target, it.get("label", "Takeaway"), v)
            elif k == "chart":
                self.figure_chart(target, v, width_mm)
            elif k == "image":
                self.figure_image(target, v, width_mm)
            elif k == "table":
                self.data_table(target, v, width_mm)
            elif k == "kpis":
                self.kpis(target, v, width_mm)
            elif k == "pagebreak":
                target.add_paragraph().add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)
            elif k in ("note", "label"):
                continue  # margin items handled by narrative()
            else:
                raise ValueError("unknown flow item: " + k)

    # ----- layouts
    def cover(self, b):
        d = self.doc
        mt = self.meta
        self.section("cover", None, margins={"top": 0, "bottom": 0, "left": 0, "right": 0})
        variant = b.get("variant", "slab")
        kicker = (b.get("kicker") or mt.get("type_label", "Report")).upper()
        if mt.get("illustrative"):
            kicker += "  \u00b7  ILLUSTRATIVE SAMPLE"
        if variant == "slab":
            t = make_table(d, 3, [34, PW - 34])
            slab = t.cell(0, 0).merge(t.cell(2, 0))
            shade(slab, C["ink"])
            text_direction(slab)
            pad(slab, 18, 0, 14, 0)
            valign(slab, "center")
            para(slab, mt.get("date", "") + "    " + mt.get("client", ""), "SlabLabel")
            for r, h in zip(t.rows, (60, 155, 80)):
                set_row_height(r, h, exact=True)
            c0 = t.cell(0, 1)
            pad(c0, 22, 14, 0, 18)
            para(c0, kicker, "Kicker")
            c1 = t.cell(1, 1)
            pad(c1, 0, 14, 4, 18)
            valign(c1, "bottom")
            rule_para(c1, PW - 34 - 14 - 18, 24, after=16)
            para(c1, mt["title"], "Display")
            if mt.get("subtitle"):
                para(c1, mt["subtitle"], "Lead", space_before=Pt(4)).paragraph_format.right_indent = Mm(6)
            c2 = t.cell(2, 1)
            pad(c2, 0, 14, 16, 18)
            valign(c2, "bottom")
            self._cover_meta(c2, PW - 34 - 14 - 18)
        else:  # band: white upper field, ink band at the bottom carrying the metadata
            t = make_table(d, 2, [PW])
            for r, h in zip(t.rows, (205, 90)):
                set_row_height(r, h, exact=True)
            c0 = t.cell(0, 0)
            pad(c0, 26, 22, 0, 22)
            para(c0, kicker, "Kicker")
            rule_para(c0, PW - 44, 24, after=16, before=70)
            para(c0, mt["title"], "Display")
            if mt.get("subtitle"):
                para(c0, mt["subtitle"], "Lead").paragraph_format.right_indent = Mm(40)
            c1 = t.cell(1, 0)
            shade(c1, C["ink"])
            pad(c1, 14, 22, 12, 22)
            valign(c1, "bottom")
            self._cover_meta(c1, PW - 44, light=True)

    def _cover_meta(self, cell, width_mm, light=False):
        mt = self.meta
        pairs = [(k, mt[v]) for k, v in (("Prepared for", "client"), ("Prepared by", "author"), ("Date", "date"), ("Version", "version")) if mt.get(v)]
        if not pairs:
            return
        n = len(pairs)
        nt = make_table(cell, 1, [width_mm / n] * n)
        for i, (k, v) in enumerate(pairs):
            c = nt.cell(0, i)
            borders(c, top=(T["rule_pt"]["base"], "FFFFFF" if light else C["ink"]))
            pad(c, 2.5, 0, 0, 3)
            para(c, k, "SlabLabel" if light else "Label")
            para(c, v, "SlabText" if light else "Meta")
        tiny(cell.add_paragraph())

    def contents(self, b):
        d = self.doc
        self.section("contents", "Contents")
        left, right = 52, CW - 52
        t = make_table(d, 1, [left, right])
        a, c = t.cell(0, 0), t.cell(0, 1)
        pad(a, 0, 0, 0, 6)
        pad(c, 0, 6, 0, 0)
        para(a, "Contents", "Heading 1")
        rule_para(a, left - 6)
        para(a, b.get("note", ""), "Note", space_before=Pt(10))
        pos = Mm(right - 6)
        for e in self.toc_entries:
            lvl, num, title = e["level"], e["number"], e["title"]
            pg = self.pages.get(title, "")
            if lvl == 1:
                p = para(c, "", "Toc1")
                if num:
                    r = p.add_run(num + "   ")
                    r.font.color.rgb = _rgb(C["accent"])
                p.add_run(title)
            else:
                p = para(c, title, "Toc2")
            p.paragraph_format.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            p.add_run("\t" + str(pg))
            if lvl == 1 and not num:
                p.paragraph_format.space_before = Pt(8)

    def summary(self, b):
        d = self.doc
        self.section("summary", b.get("title", "Executive summary"))
        left, right = 52, CW - 52
        t = make_table(d, 1, [left, right])
        set_row_height(t.rows[0], PH - M["top"] - M["bottom"] - 8)
        a, c = t.cell(0, 0), t.cell(0, 1)
        shade(a, C["accent"])
        pad(a, 10, 6, 8, 6)
        pad(c, 0, 10, 0, 0)
        para(a, b.get("slab_label", "At a glance"), "SlabLabel", space_after=Pt(10))
        for f in b.get("figures", []):
            p = para(a, "", "SlabNum")
            p.add_run(str(f["value"]))
            if f.get("unit"):
                p.add_run(f["unit"]).font.size = Pt(14)
            para(a, f["label"], "SlabText", self.cites)
        self.heading_in(c, b.get("title", "Executive summary"), b.get("kicker", "Summary"), CW - 52 - 10)
        para(c, b["lead"], "Lead", self.cites)
        for i, pt in enumerate(b.get("points", []), 1):
            p = para(c, "", "Normal", space_before=Pt(4))
            r = p.add_run(f"{i:02d}")
            r.font.name = SANS; r.bold = True; r.font.size = Pt(8); r.font.color.rgb = _rgb(C["accent"])
            p.add_run("\t")
            runs(p, pt, self.cites)
            p.paragraph_format.left_indent = Mm(10)
            p.paragraph_format.first_line_indent = Mm(-10)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(10))
            p_border(p, top=(T["rule_pt"]["hair"], C["rule"], 6))
        for it in b.get("after", []):
            self.flow(c, [it], right - 10)

    def heading_in(self, cell, title, kicker, width):
        para(cell, kicker, "Kicker")
        para(cell, title, "Heading 1")
        rule_para(cell, width)

    def opener(self, b):
        d = self.doc
        self.section("opener", None, frame=True)
        t = make_table(d, 3, [14, CW - 14])
        lab = t.cell(0, 0).merge(t.cell(2, 0))
        text_direction(lab)
        valign(lab, "bottom")
        pad(lab, 0, 0, 0, 0)
        para(lab, b.get("label", "Section " + b["number"]), "Label")
        for r, h, exact in zip(t.rows, (95, 70, 70), (True, True, True)):
            set_row_height(r, h, exact=exact)
        n = t.cell(0, 1)
        pad(n, 8, 4, 0, 0)
        para(n, b["number"], "NumXL")
        ti = t.cell(1, 1)
        valign(ti, "bottom")
        pad(ti, 0, 4, 6, 0)
        rule_para(ti, CW - 14 - 4)
        para(ti, b["title"], "Display")
        it = t.cell(2, 1)
        pad(it, 6, 4, 0, 30)
        if b.get("intro"):
            para(it, b["intro"], "Lead", self.cites)

    def narrative(self, b):
        d = self.doc
        title = b["title"]
        items = b["items"]
        layout = b.get("layout") or ("margin" if any(isinstance(i, dict) and next(iter(i)) in ("note", "quote") for i in items) else "measure")
        if layout == "twocol":
            self.section("narrative-2col", title)
            self.heading(title, b.get("kicker"), b.get("lead"))
            self.section("narrative-2col", title, cols=2, continuous=True)
            self.flow(d, items, (CW - 8) / 2)
            return
        self.section("narrative-" + layout, title)
        self.heading(title, b.get("kicker"), b.get("lead"), width=CW)
        if layout == "measure":
            tm = T["measure_mm"]["text"] + 12
            body = d
            sink = _Measure(d, CW - tm)
            self.flow(sink, items, tm)
        else:
            tw, mw = CW - T["measure_mm"]["margin_note"], T["measure_mm"]["margin_note"]
            rows, cur = [], None
            for it in items:
                k = next(iter(it)) if isinstance(it, dict) else "p"
                if cur is None or k == "h":
                    if cur is not None and cur["main"] and cur["main"][-1] != it and k == "h":
                        pass
                    cur = {"main": [], "side": []}
                    rows.append(cur)
                (cur["side"] if k in ("note", "quote") else cur["main"]).append(it)
            for rw in rows:
                t = make_table(d, 1, [tw, mw])
                a, m = t.cell(0, 0), t.cell(0, 1)
                pad(a, 0, 0, 0, 7)
                pad(m, 1.2, 5, 0, 0)
                borders(m, left=(T["rule_pt"]["hair"], C["rule"]))
                self.flow(a, rw["main"], tw - 7)
                for s in rw["side"]:
                    k = next(iter(s))
                    if k == "note":
                        v = s["note"]
                        if s.get("label"):
                            para(m, s["label"], "Label")
                        para(m, v, "Note", self.cites)
                    else:
                        p = para(m, s["quote"], "Quote", self.cites)
                        p_border(p, top=(T["rule_pt"]["base"], C["accent"], 8))
                        if s.get("by"):
                            para(m, s["by"], "Label")
                tiny(d.add_paragraph())

    def data(self, b):
        d = self.doc
        wide = b.get("landscape", False)
        title = b["title"]
        width = (PH - M["left"] - M["right"]) if wide else CW
        self.section("data-wide" if wide else "data", title, landscape=wide,
                     margins={**M, "top": 20, "bottom": 16} if wide else None)
        self.heading(title, b.get("kicker", "Analysis"), b.get("lead"), width=width)
        if b.get("kpis"):
            self.kpis(d, b["kpis"], width)
        self.flow(d, b.get("items", []), width)
        if b.get("takeaway"):
            self.callout(d, b.get("takeaway_label", "What this means"), b["takeaway"])

    def findings(self, b):
        d = self.doc
        title = b["title"]
        self.section("findings", title)
        self.heading(title, b.get("kicker", "Findings"), b.get("lead"))
        a, m = 20, 46
        mid = CW - a - m
        items = b["items"]
        t = make_table(d, len(items), [a, mid, m])
        for i, it in enumerate(items):
            row_flags(t.rows[i])
            top = (T["rule_pt"]["base"], C["ink"]) if i == 0 else (T["rule_pt"]["hair"], C["rule"])
            cells = [t.cell(i, j) for j in range(3)]
            for c in cells:
                borders(c, top=top, bottom=(T["rule_pt"]["base"], C["ink"]) if i == len(items) - 1 else None)
            pad(cells[0], 3, 0, 3, 0); pad(cells[1], 3, 0, 3, 8); pad(cells[2], 3.4, 0, 3, 0)
            hot = str(it.get("meta", {}).get("Priority", "")).lower() == "high"
            p = para(cells[0], f"{i + 1:02d}", "NumL")
            if not hot:
                for r in p.runs:
                    r.font.color.rgb = _rgb(C["ink"])
            para(cells[1], it["title"], "Heading 2", space_before=Pt(0), keep_with_next=False)
            for chunk in (it["body"] if isinstance(it["body"], list) else [it["body"]]):
                para(cells[1], chunk, "Normal", self.cites)
            for k, v in it.get("meta", {}).items():
                para(cells[2], k, "Label", space_after=Pt(0))
                para(cells[2], str(v), "Meta", space_after=Pt(4))

    def references(self, b):
        d = self.doc
        title = b.get("title", "References")
        self.section("references", title)
        self.heading(title, b.get("kicker", "Sources"), b.get("lead"))
        self.section("references", title, cols=2, continuous=True)
        for i, r in enumerate(b["items"], 1):
            self.refs.add(i)
            p = para(d, f"{i}\t" + r, "Ref", self.cites)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(6))

    def appendix(self, b):
        d = self.doc
        wide = b.get("landscape", False)
        width = (PH - M["left"] - M["right"]) if wide else CW
        self.section("appendix-wide" if wide else "appendix", b["title"], landscape=wide,
                     margins={**M, "top": 20, "bottom": 16} if wide else None)
        self.heading(b["title"], b.get("kicker", "Appendix"), b.get("lead"), width=width)
        self.flow(d, b.get("items", []), width)

    def closing(self, b):
        d = self.doc
        mt = self.meta
        self.section("closing", None, margins={"top": 0, "bottom": 0, "left": 0, "right": 0})
        t = make_table(d, 2, [PW])
        for r, h in zip(t.rows, (120, 165)):
            set_row_height(r, h, exact=True)
        top, bot = t.cell(0, 0), t.cell(1, 0)
        shade(top, C["ink"])
        pad(top, 0, 22, 18, 22)
        valign(top, "bottom")
        hp = para(top, b.get("headline", ""), "Display")
        for r in hp.runs:
            r.font.color.rgb = _rgb("FFFFFF")
        pad(bot, 14, 22, 0, 22)
        for h, txt in b.get("notes", []):
            para(bot, h, "Label")
            para(bot, txt, "Meta", self.cites, space_after=Pt(10)).paragraph_format.right_indent = Mm(50)

    # ----- build
    def plan_toc(self, blocks):
        self.toc_entries = []
        have_opener = False
        for b in blocks:
            if b["type"] == "opener":
                self.toc_entries.append({"level": 1, "number": b["number"], "title": b["title"]})
                have_opener = True
            elif b["type"] in ("narrative", "data", "findings", "appendix", "references", "summary") and b.get("toc", True):
                title = b.get("title", "Executive summary" if b["type"] == "summary" else "References")
                lvl = 2 if (have_opener and b["type"] in ("narrative", "data", "findings")) else 1
                self.toc_entries.append({"level": lvl, "number": "", "title": title})
                if b["type"] in ("references", "appendix"):
                    have_opener = False

    def build(self, blocks, out):
        self.plan_toc(blocks)
        for b in blocks:
            getattr(self, {"summary": "summary", "opener": "opener"}.get(b["type"], b["type"]))(b)
        for p in self.doc.element.body.iter(qn("w:p")):
            ppr = p.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None and not p.findall(qn("w:r")):
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
                sp.set(qn("w:line"), "20"); sp.set(qn("w:lineRule"), "exact")
                ppr.insert(0, sp)
        missing = sorted(self.cites - self.refs) if self.refs else []
        if missing:
            print("WARNING: citations without a reference entry:", missing)
        tiny(self.doc.add_paragraph())
        fix_order(self.doc.element)
        fix_order(self.doc.styles.element)
        for sec in self.doc.sections:
            fix_order(sec.header._element)
            fix_order(sec.footer._element)
        self.doc.save(out)
        return self.sequence


class _Measure:
    """Wraps the document so that paragraphs get a right indent that limits line length."""
    def __init__(self, doc, right_mm):
        self._d, self._r = doc, right_mm

    @property
    def paragraphs(self):
        return self._d.paragraphs

    def add_paragraph(self, *a, **k):
        p = self._d.add_paragraph(*a, **k)
        p.paragraph_format.right_indent = Mm(self._r)
        return p

    def add_table(self, *a, **k):
        return self._d.add_table(*a, **k)
