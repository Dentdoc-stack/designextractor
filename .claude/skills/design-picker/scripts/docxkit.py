"""docxkit: python-docx helpers for the DESIGN.md report designs (tested in LibreOffice 24.2).

Page-anchored shapes and pictures, editable text boxes, exact-leading paragraphs, fixed-layout tables,
one-section-per-page layout with optional folio, Lucide icon chips. Usage:

    import docxkit as k
    k.configure("out", font="Figtree", ink="1B1D2E", soft="5A5F6E", footer="REPORT · 2025", icons="icons")
    k.page(header_shapes=[lambda r: k.shape(r, 0, 0, 70, 297, fill="1A2E5F", behind=True, z=1)])
    k.P(k.doc, "TITLE", 36, font="Figtree ExtraBold")
    k.save("report.docx")

Weights: bold = base family + bold; SemiBold / ExtraBold are their own family names ("Figtree ExtraBold").
"""
import itertools
import re
from pathlib import Path

import cairosvg
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"
CFG = dict(out=Path("out"), font="Figtree", ink="1B1D2E", soft="5A5F6E", footer="", icons=None)

# ----------------------------------------------------------------------------- low-level XML
EMU = 36000
NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"')
_ids = itertools.count(1000)


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def r_xml(text, font=None, size=10, color=None, bold=False, italic=False, track=0, sup=False):
    font, color = font or CFG["font"], color or CFG["ink"]
    rpr = (f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:cs="{font}" w:eastAsia="{font}"/>'
           f'{"<w:b/>" if bold else ""}{"<w:i/>" if italic else ""}<w:color w:val="{color}"/>'
           f'{f"<w:spacing w:val=\"{track}\"/>" if track else ""}<w:sz w:val="{int(size * 2)}"/><w:szCs w:val="{int(size * 2)}"/>'
           f'{"<w:vertAlign w:val=\"superscript\"/>" if sup else ""}')
    parts = esc(text).split("\n")
    body = '<w:br/>'.join(f'<w:t xml:space="preserve">{t}</w:t>' for t in parts)
    return f'<w:r {NS}><w:rPr>{rpr}</w:rPr>{body}</w:r>'


def p_xml(runs, align="left", before=0, after=0, line=None):
    sp = f'<w:spacing w:before="{int(before * 20)}" w:after="{int(after * 20)}"' + (
        f' w:line="{int(line * 20)}" w:lineRule="exact"/>' if line else "/>")
    jc = {"left": "left", "center": "center", "right": "right"}[align]
    return f'<w:p {NS}><w:pPr>{sp}<w:jc w:val="{jc}"/></w:pPr>{"".join(runs)}</w:p>'


def T(text, size=10, color=None, font=None, bold=False, lead=1.38, align="left", before=0, after=0, track=0, italic=False):
    """One text-box paragraph (exact leading)."""
    return p_xml([r_xml(text, font, size, color, bold, italic, track)], align, before, after, size * lead)


def _anchor_open(x, y, w, h, behind, z, name):
    i = next(_ids)
    return (f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{z}" '
            f'behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(x * EMU)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(y * EMU)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{int(w * EMU)}" cy="{int(h * EMU)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:wrapNone/><wp:docPr id="{i}" name="{name} {i}"/><wp:cNvGraphicFramePr/>')


def shape(run, x, y, w, h, fill=None, geom="rect", line=None, behind=False, z=10, text=None,
          inset=(0, 0, 0, 0), anchor="t", rot=0):
    """Page-anchored DrawingML shape; text = list of p_xml strings (editable text box)."""
    fill_x = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else '<a:noFill/>'
    ln_x = (f'<a:ln w="{int(line[0] * 12700)}"><a:solidFill><a:srgbClr val="{line[1]}"/></a:solidFill></a:ln>'
            if line else '<a:ln><a:noFill/></a:ln>')
    txbx = f'<wps:txbx><w:txbxContent>{"".join(text)}</w:txbxContent></wps:txbx>' if text else ''
    l, t, r, b = (int(v * EMU) for v in inset)
    rot_x = f' rot="{int(rot * 60000)}"' if rot else ''
    wsp = (f'<wps:wsp><wps:cNvSpPr/><wps:spPr><a:xfrm{rot_x}><a:off x="0" y="0"/><a:ext cx="{int(w * EMU)}" cy="{int(h * EMU)}"/></a:xfrm>'
           f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>{fill_x}{ln_x}</wps:spPr>{txbx}'
           f'<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="{l}" tIns="{t}" rIns="{r}" bIns="{b}" anchor="{anchor}">'
           f'<a:noAutofit/></wps:bodyPr></wps:wsp>')
    xml = (f'<mc:AlternateContent {NS}><mc:Choice Requires="wps"><w:drawing>'
           f'{_anchor_open(x, y, w, h, behind, z, "Shape")}<a:graphic><a:graphicData '
           f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">{wsp}</a:graphicData></a:graphic>'
           f'</wp:anchor></w:drawing></mc:Choice></mc:AlternateContent>')
    run._r.append(parse_xml(xml))


def picture(run, path, x, y, w, h, behind=True, z=2):
    """Page-anchored picture (inline picture moved into a wp:anchor)."""
    inline = run.part.new_pic_inline(str(path), Mm(w), Mm(h))
    graphic = inline.find(qn("a:graphic"))
    xml = f'<w:drawing {NS}>{_anchor_open(x, y, w, h, behind, z, "Picture")}</wp:anchor></w:drawing>'
    drawing = parse_xml(xml)
    drawing[0].append(graphic)
    run._r.append(drawing)


def text_box(run, x, y, w, h, paras, fill=None, anchor="t", inset=(0, 0, 0, 0), z=20, line=None, geom="rect"):
    shape(run, x, y, w, h, fill=fill, text=paras, anchor=anchor, inset=inset, z=z, line=line, geom=geom)


# ----------------------------------------------------------------------------- document helpers
doc = None
_first_section = [True]


def configure(out, font="Figtree", ink="1B1D2E", soft="5A5F6E", footer="", icons=None):
    """Create the document and set defaults. Call once before page()."""
    global doc
    CFG.update(out=Path(out), font=font, ink=ink, soft=soft, footer=footer, icons=Path(icons) if icons else None)
    CFG["out"].mkdir(parents=True, exist_ok=True)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = font, Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.paragraph_format.space_after = Pt(0)
    _first_section[0] = True
    for f in FONTS.glob(font.replace(" ", "") + "-*.ttf"):
        font_manager.fontManager.addfont(str(f))
    plt.rcParams["font.family"] = font
    return doc


def save(name):
    out = CFG["out"] / name
    doc.save(out)
    return out


def _folio(fp, left, right):
    width = 210 - left - right
    fp.paragraph_format.line_spacing = Pt(10)
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(width), WD_TAB_ALIGNMENT.RIGHT)
    fp._p.append(parse_xml(r_xml(CFG["footer"], CFG["font"], 6.5, CFG["soft"], track=20)))
    fp._p.append(parse_xml(f'<w:r {NS}><w:tab/></w:r>'))
    fp._p.append(parse_xml(f'<w:fldSimple {NS} w:instr="PAGE">{r_xml("1", CFG["font"] + " SemiBold", 7.5, CFG["soft"])}</w:fldSimple>'))


def _reset(hf):
    hf.is_linked_to_previous = False
    p = hf.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(1)
    return p


def page(top=24, bottom=20, left=22, right=20, folio=True, header_shapes=(), first_header_shapes=None):
    """Start a new section on a new page. header_shapes: callables(run) drawn behind text on every page of
    the section; first_header_shapes: drawn on the section's first page only (its own header)."""
    if _first_section[0]:
        sec = doc.sections[0]
        _first_section[0] = False
    else:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(top), Mm(bottom), Mm(left), Mm(right)
    sec.header_distance, sec.footer_distance = Mm(6), Mm(10)
    sec.different_first_page_header_footer = first_header_shapes is not None
    pairs = [(sec.header, sec.footer, header_shapes)]
    if first_header_shapes is not None:
        pairs.append((sec.first_page_header, sec.first_page_footer, first_header_shapes))
    for hdr, ftr, shapes in pairs:
        hp, fp = _reset(hdr), _reset(ftr)
        for draw in shapes:
            draw(hp.add_run())
        if folio:
            _folio(fp, left, right)
    return sec


def P(container, text="", size=10, color=None, font=None, bold=False, lead=1.5, before=0, after=0, align="left",
      track=0, right_indent=0, left_indent=0, keep=False, runs=None, italic=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(size * lead)
    pf.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if right_indent:
        pf.right_indent = Mm(right_indent)
    if left_indent:
        pf.left_indent = Mm(left_indent)
    pf.keep_with_next = keep
    for rx in (runs or [r_xml(text, font, size, color, bold, italic, track)]):
        if text or runs:
            p._p.append(parse_xml(rx))
    return p


def cell_p(cell, *a, **k):
    """Paragraph in a table cell; the cell's initial empty paragraph is reused once."""
    if not getattr(cell, "_used", False) and len(cell.paragraphs) == 1 and not cell.paragraphs[0].runs:
        cell._used = True
        cell._element.remove(cell.paragraphs[0]._p)
    return P(cell, *a, **k)


def _border_xml(tag, edges):
    return f'<w:{tag} {NS}>' + "".join(
        f'<w:{e} w:val="{"single" if v else "nil"}" w:sz="{int(v[0] * 8) if v else 0}" w:space="0" w:color="{v[1] if v else "auto"}"/>'
        for e, v in edges.items()) + f'</w:{tag}>'


def table(container, widths, rows, cell_margin=(0, 0, 0, 0)):
    t = container.add_table(rows=rows, cols=len(widths))
    tpr = t._tbl.tblPr
    keep = [el for el in tpr if el.tag in (qn("w:tblStyle"),)]
    for el in list(tpr):
        tpr.remove(el)
    for el in keep:
        tpr.append(el)
    tpr.append(parse_xml(f'<w:tblW {NS} w:w="{int(sum(widths) * 56.693)}" w:type="dxa"/>'))
    tpr.append(parse_xml(_border_xml("tblBorders", {e: None for e in ("top", "left", "bottom", "right", "insideH", "insideV")})))
    tpr.append(parse_xml(f'<w:tblLayout {NS} w:type="fixed"/>'))
    l, t_, r, b = cell_margin
    tpr.append(parse_xml(f'<w:tblCellMar {NS}><w:top w:w="{int(t_ * 56.693)}" w:type="dxa"/><w:left w:w="{int(l * 56.693)}" w:type="dxa"/>'
                         f'<w:bottom w:w="{int(b * 56.693)}" w:type="dxa"/><w:right w:w="{int(r * 56.693)}" w:type="dxa"/></w:tblCellMar>'))
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 56.693)))
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Mm(w)
    return t


def cell_style(cell, fill=None, borders=None, valign=None):
    tcpr = cell._tc.get_or_add_tcPr()
    if borders:
        tcpr.append(parse_xml(_border_xml("tcBorders", borders)))
    if fill:
        tcpr.append(parse_xml(f'<w:shd {NS} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
    if valign:
        tcpr.append(parse_xml(f'<w:vAlign {NS} w:val="{valign}"/>'))


def row_no_split(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(parse_xml(f'<w:cantSplit {NS}/>'))
    if header:
        trpr.append(parse_xml(f'<w:tblHeader {NS}/>'))


def inline_pic(cell_or_par, path, w_mm, align="center", before=0, after=0):
    p = cell_p(cell_or_par, "", before=before, after=after, align=align, lead=1.0, size=w_mm * 2.83) \
        if hasattr(cell_or_par, "_tc") else P(cell_or_par, "", align=align, before=before, after=after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(path), width=Mm(w_mm))
    return p
def icon(name, fill, size_mm=14, fg="#FFFFFF", stroke=1.6, pad=0.27):
    """Lucide line glyph on a solid circle, 300 dpi PNG."""
    out = CFG["out"] / f"icon-{name}-{fill}-{size_mm}.png"
    inner = re.search(r"<svg[^>]*>(.*)</svg>", (CFG["icons"] / f"{name}.svg").read_text(), re.S).group(1)
    s, off = 24 * (1 - 2 * pad), 24 * pad
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" width="24" height="24">'
           f'<circle cx="12" cy="12" r="12" fill="#{fill}"/><g transform="translate({off} {off}) scale({s / 24})" '
           f'fill="none" stroke="{fg}" stroke-width="{stroke}" stroke-linecap="round" stroke-linejoin="round">{inner}</g></svg>')
    px = round(size_mm / 25.4 * 300)
    cairosvg.svg2png(bytestring=svg.encode(), write_to=str(out), output_width=px, output_height=px)
    return out


def hexc(h):
    return "#" + h

